#!/usr/bin/env python3
"""
EA3 - Verificacion INDEPENDIENTE de los conteos del reporte.
No usa pandas ni la logica de merge de enrichement.py: recalcula con la libreria estandar
(sqlite3, json, csv, xml.etree, html.parser, re) + openpyxl, lee el Excel enriquecido que produjo enrichement.py
y compara con las cifras impresas en src/static/auditoria/enriched_report.txt.
Termina con codigo 1 si algo no coincide.   Uso: python src/verificar_conteos.py
"""
from __future__ import annotations

import csv
import json
import re
import sqlite3
import sys
import unicodedata
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

B = Path(__file__).resolve().parent.parent
F = B / "src" / "static" / "fuentes"
REPORT = B / "src" / "static" / "auditoria" / "enriched_report.txt"
ENR = B / "src" / "static" / "xlsx" / "enriched_data.xlsx"      # salida principal (src/xlsx/ es copia idéntica)
ENR_COPIA = B / "src" / "xlsx" / "enriched_data.xlsx"
CLEAN = B / "src" / "static" / "xlsx" / "cleaned_data.xlsx"


def n(v):
    s = unicodedata.normalize("NFKD", "" if v is None else str(v)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9]+", " ", s).strip()
    return re.sub(r"\s+region$", "", s)


def filas_xlsx(path, hoja=None):
    ws = load_workbook(path, read_only=True)[hoja] if hoja else load_workbook(path, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    cab = list(next(it))
    return cab, [dict(zip(cab, r)) for r in it]


class Tabla(HTMLParser):
    def __init__(self):
        super().__init__()
        self.filas, self._f, self._c, self.en = [], None, None, False
    def handle_starttag(self, t, a):
        if t == "tr": self._f = []
        if t in ("td", "th"): self._c, self.en = "", True
    def handle_data(self, d):
        if self.en: self._c += d
    def handle_endtag(self, t):
        if t in ("td", "th"): self._f.append(self._c.strip()); self.en = False
        if t == "tr" and self._f: self.filas.append(self._f)


def main() -> int:
    # ---- base: cleaned_data.xlsx y SQLite ----
    _, limpio = filas_xlsx(CLEAN)
    con = sqlite3.connect(B / "src" / "static" / "db" / "ingestion.db")
    db_n = con.execute("select count(*) from cerveceria").fetchone()[0]
    ids_db = {r[0] for r in con.execute("select id from cerveceria")}
    con.close()
    N = len(limpio)
    esp = {"db": db_n, "base": N, "ids_en_db": sum(r["id"] in ids_db for r in limpio)}

    # ---- fuentes con libreria estandar ----
    js = json.loads((F / "paises.json").read_text(encoding="utf-8"))
    j_pais = {n(nm): (p["iso2"]) for p in js for nm in p["nombres_dataset"]}
    _, xl = filas_xlsx(F / "paises_capital.xlsx", "paises")
    x_pais = {n(r["Pais"]) for r in xl}
    with open(F / "tipos_cerveceria.csv", encoding="utf-8", newline="") as fh:
        c_tipo = {n(r["tipo"]) for r in csv.DictReader(fh)}
    xml_k = set()
    for s in ET.parse(F / "estados.xml").getroot().iter("subdivision"):
        i2 = s.findtext("iso2")
        xml_k.add(i2 + "|" + n(s.findtext("nombre")))
        xml_k.add(i2 + "|" + n(s.findtext("abreviatura")))
    t = Tabla(); t.feed((F / "ciudades.html").read_text(encoding="utf-8"))
    cab = t.filas[0]; h_ci = {n(r[0]) + "|" + n(r[1]) + "|" + n(r[2]) for r in t.filas[1:] if len(r) == len(cab)}
    t_iso = set()
    for ln in (F / "idiomas.txt").read_text(encoding="utf-8").splitlines():
        if ln.startswith("#") or ln.startswith("iso2|") or not ln.strip(): continue
        t_iso.add(ln.split("|")[0].strip().upper())

    m = dict.fromkeys(["json", "xlsx", "csv", "xml", "html", "txt"], 0)
    completos = 0
    for r in limpio:
        p, i2 = n(r["pais"]), j_pais.get(n(r["pais"]), "")
        ok = {
            "json": p in j_pais, "xlsx": p in x_pais, "csv": n(r["tipo_cerveza"]) in c_tipo,
            "xml": (i2 + "|" + n(r["estado_provincia"])) in xml_k,
            "html": (n(r["ciudad"]) + "|" + n(r["estado_provincia"]) + "|" + p) in h_ci,
            "txt": i2.upper() in t_iso,
        }
        for k, v in ok.items(): m[k] += v
        completos += all(ok.values())
    esp["m"], esp["completos"] = m, completos

    # ---- Excel enriquecido que produjo enrichement.py ----
    cab_e, enr = filas_xlsx(ENR, "enriquecido")
    real_xlsx = {"filas": len(enr), "ids_unicos": len({r["id"] for r in enr}),
                 "iso2": sum(r["iso2"] not in (None, "") for r in enr),
                 "capital": sum(r["capital_pais"] not in (None, "") for r in enr),
                 "tipo": sum(r["descripcion_tipo"] not in (None, "") for r in enr),
                 "xml": sum(r["subdivision_iso3166_2"] not in (None, "") for r in enr),
                 "html": sum(r["zona_horaria"] not in (None, "") for r in enr),
                 "txt": sum(r["idioma_principal"] not in (None, "") for r in enr),
                 "completos": sum(str(r["enriquecimiento_completo"]).lower() == "true" for r in enr),
                 "iso_num_3d": sum(len(str(r["iso_numerico"])) == 3 for r in enr)}

    # ---- texto del reporte ----
    rep = REPORT.read_text(encoding="utf-8")
    def g(pat):
        mm = re.search(pat, rep); return int(mm.group(1)) if mm else None
    coin = [int(x) for x in re.findall(r"(?<!NO )COINCIDENTES\s+: (\d+) de", rep)]
    nocoin = [int(x) for x in re.findall(r"NO COINCIDENTES\s+: (\d+) de", rep)]
    rep_v = {"db": g(r"tabla cerveceria, sin limpiar\): (\d+) registros"),
             "base": g(r"BASE del enriquecimiento: (\d+) registros"),
             "ids_en_db": g(r"presentes en la base SQLite: (\d+) de"),
             "enriquecido": g(r"Dataset enriquecido: (\d+) registros"),
             "completos": g(r"enriquecimiento_completo=True\): (\d+) \(")}

    checks, fallos = [], 0
    def chk(nombre, a, b):
        nonlocal fallos
        ok = a == b; fallos += (not ok)
        checks.append(f"{'OK ' if ok else 'ERR'} {nombre}: independiente={a} | comparado={b}")

    chk("registros SQLite (reporte)", esp["db"], rep_v["db"])
    chk("registros base limpia (reporte)", esp["base"], rep_v["base"])
    chk("ids limpios en SQLite (reporte)", esp["ids_en_db"], rep_v["ids_en_db"])
    chk("registros enriquecidos (reporte) = base", esp["base"], rep_v["enriquecido"])
    chk("filas hoja 'enriquecido' en Excel = base", esp["base"], real_xlsx["filas"])
    chk("ids unicos en Excel = base", esp["base"], real_xlsx["ids_unicos"])
    orden = ["json", "xlsx", "csv", "xml", "html", "txt"]
    xl_key = {"json": "iso2", "xlsx": "capital", "csv": "tipo", "xml": "xml", "html": "html", "txt": "txt"}
    for i, k in enumerate(orden):
        chk(f"coincidentes {k} (reporte)", m[k], coin[i])
        chk(f"no coincidentes {k} (reporte)", N - m[k], nocoin[i])
        chk(f"coincidentes {k} (columnas del Excel)", m[k], real_xlsx[xl_key[k]])
    chk("iso_numerico con 3 digitos (ceros a la izquierda conservados)", N, real_xlsx["iso_num_3d"])
    chk("enriquecimiento_completo (reporte)", esp["completos"], rep_v["completos"])
    chk("enriquecimiento_completo (Excel)", esp["completos"], real_xlsx["completos"])

    if ENR_COPIA.exists():
        chk("src/xlsx/enriched_data.xlsx identico a src/static/xlsx/enriched_data.xlsx",
            ENR.read_bytes() == ENR_COPIA.read_bytes(), True)

    print("\n".join(checks))
    print(f"\nRESUMEN: base={N} | SQLite={db_n} | coincidentes={m} | completos={completos}")
    print("RESULTADO:", "TODO COINCIDE" if not fallos else f"{fallos} DIFERENCIAS")
    return 1 if fallos else 0


if __name__ == "__main__":
    sys.exit(main())
